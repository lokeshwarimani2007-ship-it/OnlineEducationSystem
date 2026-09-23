import docx

def verify():
    doc = docx.Document('Online_Examination_System_Project_Report.docx')
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")

    required_headings = [
        "Online Examination System using Python",
        "M.LOGESHWARI",
        "3-B.Sc Computer Science",
        "INDEX",
        "1. INTRODUCTION",
        "1.1 SYSTEM SPECIFICATION",
        "1.1.1 HARDWARE CONFIGURATION",
        "1.1.2 SOFTWARE CONFIGURATION",
        "2. SYSTEM STUDY",
        "2.1 EXISTING SYSTEM",
        "2.1.1 DESCRIPTION",
        "2.1.2 DRAWBACKS",
        "2.2 PROPOSED SYSTEM",
        "2.2.1 DESCRIPTION",
        "2.2.2 FEATURES",
        "3. SYSTEM DESIGN AND DEVELOPMENT",
        "3.1 FILE DESIGN",
        "3.2 INPUT DESIGN",
        "3.3 OUTPUT DESIGN",
        "3.4 DATABASE DESIGN",
        "3.5 SYSTEM DEVELOPMENT",
        "3.6 DESCRIPTION OF MODULES",
        "4. TESTING AND IMPLEMENTATION",
        "4.1 TEST CASES EXAMPLE",
        "4.2 IMPLEMENTATION",
        "4.3 DEPLOYMENT OPTIONS",
        "CONCLUSION",
        "BIBLIOGRAPHY",
        "APPENDICES",
        "A) SAMPLE CODING",
        "B) SAMPLE INPUT",
        "C) SAMPLE OUTPUT"
    ]

    all_text = "\n".join(p.text for p in doc.paragraphs)
    missing = [h for h in required_headings if h not in all_text]
    if missing:
        print(f"MISSING HEADINGS / SECTIONS: {missing}")
    else:
        print("\nALL REQUIRED HEADINGS AND SECTIONS ARE FULLY PRESENT AND VERIFIED!")

if __name__ == '__main__':
    verify()
